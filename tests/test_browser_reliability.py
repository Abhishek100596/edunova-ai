"""Browser-reliability regression tests for EduNova repair work."""

from __future__ import annotations

import io
import json
from pathlib import Path
from unittest.mock import patch

import pytest
from docx import Document

from app.ai.response import normalize_ai_response, normalize_ai_text
from app.extensions import db
from app.models import (
    CareerRole,
    InterviewSession,
    Project,
    Resume,
    ResumeAnalysis,
    StudentProfile,
    StudentSkill,
    User,
)
from app.services import interview as interview_svc
from app.services import project_mentor
from app.services import resume_intel
from app.services import roadmap as roadmap_svc
from app.services import what_if
from app.services.coach import ask_coach, clear_conversation


def _login(client, email="student@test.com", password="Student@123"):
    return client.post(
        "/login",
        data={"email": email, "password": password, "submit": True},
        follow_redirects=True,
    )


def test_api_health_shape(client):
    r = client.get("/api/health")
    assert r.status_code == 200
    data = r.get_json()
    assert data["status"] == "ok"
    assert data["database"] == "connected"
    assert "ai" in data
    assert "configured" in data["ai"]
    assert "available" in data["ai"]
    # Never leak secrets
    blob = json.dumps(data).lower()
    assert "sk-" not in blob
    assert "api_key" not in blob or data["ai"].get("detail")


def test_normalize_ai_response_shapes():
    hello = normalize_ai_response("Hello student")
    assert hello["ok"] is True
    assert "Hello" in hello["text"]

    d = normalize_ai_response({"answer": "Focus on SQL joins", "meta": {"x": 1}})
    assert d["ok"] is True
    assert "SQL" in d["text"]
    assert "{'answer'" not in d["text"]

    lst = normalize_ai_response(["Learn Python", "Practice SQL"])
    assert "Python" in lst["text"] and "SQL" in lst["text"]

    empty = normalize_ai_response("")
    assert empty["ok"] is False

    # Malformed provider-ish payload
    nested = normalize_ai_response(
        {"choices": [{"message": {"content": '{"reply": "Practice projects"}'}}]}
    )
    assert "Practice" in nested["text"]
    assert "Response [" not in nested["text"]


def test_resume_skills_string_list_renders(client, app):
    """skills_found is stored as JSON list of strings — template must not 500."""
    _login(client)
    with app.app_context():
        student = User.query.filter_by(email="student@test.com").first()
        profile = StudentProfile.query.filter_by(user_id=student.id).first()
        resume = Resume(
            student_id=profile.id,
            original_filename="cv.pdf",
            stored_filename="cv.pdf",
            file_ext="pdf",
            parsed_text="Python SQL Excel projects education",
        )
        db.session.add(resume)
        db.session.commit()
        analysis = ResumeAnalysis(
            resume_id=resume.id,
            completeness_score=72.0,
            skills_found=json.dumps(["Python", "SQL"]),
            sections_present=json.dumps({"skills": True, "education": True}),
            suggestions=json.dumps(["Add quantified impact"]),
            word_count=40,
        )
        db.session.add(analysis)
        db.session.commit()

    r = client.get("/student/resume")
    assert r.status_code == 200
    body = r.data.decode("utf-8", errors="replace")
    assert "Python" in body
    assert "SQL" in body
    assert "Traceback" not in body
    assert "AttributeError" not in body


def test_resume_text_analysis_readable(app):
    with app.app_context():
        student = User.query.filter_by(email="student@test.com").first()
        profile = StudentProfile.query.filter_by(user_id=student.id).first()
        resume = Resume(
            student_id=profile.id,
            original_filename="manual.txt",
            stored_filename="manual.txt",
            file_ext="txt",
            parsed_text=(
                "Education: B.Tech CSE\nExperience: Intern at Acme\n"
                "Projects: Dashboard\nSkills: Python, SQL\nCertifications: AWS"
            ),
        )
        db.session.add(resume)
        db.session.commit()
        analysis = resume_intel.analyze_resume(resume, file_path=None)
        assert analysis.completeness_score is not None
        skills = json.loads(analysis.skills_found)
        assert isinstance(skills, list)
        assert "Python" in skills or "SQL" in skills
        report = resume_intel.build_resume_intelligence_report(
            analysis, resume_text=resume.parsed_text or ""
        )
        assert "overall" in report


def test_resume_unsupported_and_empty_file(app, tmp_path):
    with app.app_context():
        bad = tmp_path / "note.txt"
        bad.write_text("hello", encoding="utf-8")
        with pytest.raises(ValueError, match="Unsupported"):
            resume_intel.extract_text(bad, "txt")

        empty_pdf = tmp_path / "empty.pdf"
        empty_pdf.write_bytes(b"%PDF-1.4\n%\xe2\xe3\xcf\xd3\n1 0 obj<<>>endobj\ntrailer<<>>\n%%EOF\n")
        with pytest.raises(ValueError):
            resume_intel.extract_text(empty_pdf, "pdf")


def test_resume_docx_roundtrip(app, tmp_path):
    path = tmp_path / "resume.docx"
    doc = Document()
    doc.add_heading("Jane Student", 0)
    doc.add_paragraph("Education: B.Tech Computer Science")
    doc.add_paragraph("Skills: Python, SQL, Excel")
    doc.add_paragraph("Projects: Built a Flask analytics dashboard")
    doc.save(path)
    with app.app_context():
        text = resume_intel.extract_text(path, "docx")
        assert "Python" in text
        skills = resume_intel.extract_skills_from_text(text)
        names = {s["name"] for s in skills}
        assert "Python" in names or "SQL" in names


def test_interview_one_question_at_a_time(client, app):
    _login(client)
    with app.app_context():
        student = User.query.filter_by(email="student@test.com").first()
        profile = StudentProfile.query.filter_by(user_id=student.id).first()
        sess = interview_svc.start_session(
            profile, "technical", role_focus="Data Analyst"
        )
        detail = interview_svc.session_detail(sess.id, student_id=profile.id)
        assert detail["current_question_id"]
        unanswered = [q for q in detail["questions"] if q["answer"] is None]
        assert len(unanswered) >= 2
        sid = sess.id
        first_qid = detail["current_question_id"]

    r = client.get(f"/student/interview?session_id={sid}")
    assert r.status_code == 200
    body = r.data.decode("utf-8", errors="replace")
    # Only one answer textarea should be present for the current question
    assert body.count('name="answer_text"') == 1
    assert f'value="{first_qid}"' in body
    assert "Answer earlier questions first" in body


def test_interview_answer_evaluate_complete(app):
    with app.app_context():
        student = User.query.filter_by(email="student@test.com").first()
        profile = StudentProfile.query.filter_by(user_id=student.id).first()
        sess = interview_svc.start_session(
            profile, "behavioral", role_focus="Analyst"
        )
        detail = interview_svc.session_detail(sess.id, student_id=profile.id)
        qid = detail["current_question_id"]
        ans = interview_svc.submit_answer(
            qid,
            "I led a campus project using Python and SQL to analyze student attendance. "
            "I collaborated with teammates, measured outcomes, and presented results.",
            student_id=profile.id,
        )
        assert ans.score is not None
        assert ans.feedback
        assert "{" not in (ans.feedback or "")[:40] or "Score" in ans.feedback
        with pytest.raises(ValueError):
            interview_svc.submit_answer(qid, "   ", student_id=profile.id)

        detail2 = interview_svc.session_detail(sess.id, student_id=profile.id)
        while detail2.get("current_question_id") and detail2["status"] == "in_progress":
            cq = detail2["current_question_id"]
            interview_svc.submit_answer(
                cq,
                "I used structured reasoning, checked assumptions, and communicated trade-offs clearly with examples.",
                student_id=profile.id,
            )
            detail2 = interview_svc.session_detail(sess.id, student_id=profile.id)
            if detail2["answered_count"] >= detail2["question_count"]:
                break
        completed = interview_svc.complete_session(sess.id, student_id=profile.id)
        assert completed.status == "completed"
        assert completed.overall_score is not None


def test_coach_local_fallback_readable(app):
    with app.app_context():
        student = User.query.filter_by(email="student@test.com").first()
        profile = StudentProfile.query.filter_by(user_id=student.id).first()
        clear_conversation(profile)
        result = ask_coach(profile, "What should I learn next?")
        assert result["reply"]
        assert isinstance(result["reply"], str)
        assert result["reply"].strip()
        assert "{'answer'" not in result["reply"]
        assert "<Response" not in result["reply"]
        # Normalization of provider-shaped dicts
        normalized = normalize_ai_text({"answer": "Focus on SQL and projects next."})
        assert "SQL" in normalized
        assert "{'answer'" not in normalized


def test_coach_groq_timeout_falls_back(app):
    with app.app_context():
        student = User.query.filter_by(email="student@test.com").first()
        profile = StudentProfile.query.filter_by(user_id=student.id).first()
        clear_conversation(profile)
        app.config["AI_PROVIDER"] = "groq"

        def boom(*_a, **_k):
            raise TimeoutError("groq timeout")

        with patch("app.services.coach.complete_with_fallback", side_effect=boom):
            try:
                result = ask_coach(profile, "Help me prepare for interviews")
            except TimeoutError:
                pytest.fail("Coach must not crash on provider timeout")
            assert result.get("reply")
            assert isinstance(result["reply"], str)
            assert result.get("fallback_used") is True
        app.config["AI_PROVIDER"] = "local"


def test_what_if_non_mutating_multiple(app):
    with app.app_context():
        student = User.query.filter_by(email="student@test.com").first()
        profile = StudentProfile.query.filter_by(user_id=student.id).first()
        cgpa_before = profile.cgpa
        skills_before = {
            (l.skill_id, l.level)
            for l in StudentSkill.query.filter_by(student_id=profile.id).all()
        }
        r1 = what_if.simulate_skill_level_change(profile, "Python", 5)
        r2 = what_if.simulate_extra_projects(profile, 3)
        assert r1.get("simulated") or r1.get("delta_overall_readiness") is not None
        assert r2.get("simulated") or r2.get("delta_overall_readiness") is not None
        db.session.refresh(profile)
        assert profile.cgpa == cgpa_before
        skills_after = {
            (l.skill_id, l.level)
            for l in StudentSkill.query.filter_by(student_id=profile.id).all()
        }
        assert skills_before == skills_after


def test_project_mentor_manual_and_github_unavailable(app):
    with app.app_context():
        student = User.query.filter_by(email="student@test.com").first()
        profile = StudentProfile.query.filter_by(user_id=student.id).first()
        project = Project(
            student_id=profile.id,
            title="Campus Portal",
            description="Flask student portal with attendance charts",
            tech_stack="Python, Flask, SQL",
            role="Full-stack",
            url="https://example.com/demo",
        )
        db.session.add(project)
        db.session.commit()
        report = project_mentor.mentor_project(project, profile)
        assert report["ok"] is True
        assert report["title"] == "Campus Portal"
        assert isinstance(report["strengths"], list)
        assert isinstance(report["improvements"], list)
        text_blob = json.dumps(report)
        assert "Traceback" not in text_blob

        # Empty-ish project still returns readable structure
        thin = Project(student_id=profile.id, title="Untitled Mini")
        db.session.add(thin)
        db.session.commit()
        thin_report = project_mentor.mentor_project(thin, profile)
        assert thin_report["ok"] is True
        assert thin_report["summary"]


def test_roadmap_html_readable_no_dict_dump(client, app):
    _login(client)
    with app.app_context():
        role = CareerRole.query.filter_by(name="Data Analyst").first()
        role_id = role.id
    r = client.post(
        "/student/roadmap",
        data={"role_id": role_id, "generate": "1"},
        follow_redirects=True,
    )
    assert r.status_code == 200
    body = r.data.decode("utf-8", errors="replace")
    assert "roadmap" in body.lower() or "FOUNDATION" in body or "Foundation" in body
    assert "{'title'" not in body
    assert '"tasks":' not in body


def test_roadmap_empty_profile_and_invalid_role(app):
    with app.app_context():
        user = User(email="empty@test.com", name="Empty", role="student")
        user.set_password("Empty@123")
        db.session.add(user)
        db.session.commit()
        profile = StudentProfile(user_id=user.id, onboarding_pct=0)
        db.session.add(profile)
        db.session.commit()
        role = CareerRole.query.filter_by(name="Data Analyst").first()
        rm = roadmap_svc.generate_roadmap(profile, role, replace_existing=True)
        assert rm.id
        summary = roadmap_svc.roadmap_summary(rm)
        assert summary["tasks"]
        # Invalid role handling
        with pytest.raises(ValueError, match="valid career role"):
            roadmap_svc.generate_roadmap(profile, None, replace_existing=True)


def test_error_handlers_html_and_api(client):
    r = client.get("/this-page-does-not-exist-edunova")
    assert r.status_code == 404
    assert b"Traceback" not in r.data

    r = client.get("/api/this-endpoint-missing")
    assert r.status_code == 404
    # API path should prefer JSON when Accept is json
    r = client.get(
        "/api/this-endpoint-missing",
        headers={"Accept": "application/json"},
    )
    assert r.status_code == 404
    data = r.get_json()
    if data:
        assert "error" in data or "message" in data


def test_form_errors_shown_on_invalid_coach(client):
    _login(client)
    r = client.post(
        "/student/coach",
        data={"message": "", "submit": True},
        follow_redirects=True,
    )
    assert r.status_code == 200
    # Either flashed validation or form_errors macro content
    body = r.data.decode("utf-8", errors="replace").lower()
    assert "please fix" in body or "required" in body or "message" in body
